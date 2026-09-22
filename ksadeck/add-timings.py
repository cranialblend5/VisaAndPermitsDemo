#!/usr/bin/env python3
"""Stamp a run-sheet timing onto each slide's speaker notes.

PLAN is zipped against the deck in slide order, so a label here must match the
slide at that position. Insert a slide and every label below it shifts.

The deck is built to be delivered in twenty minutes. These add up to 1200
seconds exactly, so a presenter can see at a glance whether they are ahead or
behind. Run after `node build.js`.
"""
from pptx import Presentation
import pathlib, sys

DECK = pathlib.Path(__file__).parent / "INK_IT_Visa_and_Permits_KSA.pptx"

PLAN = [
    (1,  40,  "Open"),
    (2,  80,  "Set the problem"),
    (3,  70,  "What it is"),
    (4,  85,  "Permit types"),
    (5,  95,  "The map"),
    (6,  140, "Walk the permit that hurts"),
    (7,  115, "Where the deal widens"),
    (8,  145, "Integration answer"),
    (9,  80,  "Risk reframe"),
    (10, 70,  "Who uses it"),
    (11, 65,  "Accountability and cost"),
    (12, 70,  "Technical credibility"),
    (13, 85,  "What you provision"),
    (14, 60,  "Close wide"),
]


def main() -> int:
    total = sum(sec for _, sec, _ in PLAN)
    if total != 1200:
        print(f"plan totals {total}s, not 1200", file=sys.stderr)
        return 1

    prs = Presentation(DECK)
    if len(prs.slides._sldIdLst) != len(PLAN):
        print(f"deck has {len(prs.slides._sldIdLst)} slides, plan has {len(PLAN)}", file=sys.stderr)
        return 1

    elapsed = 0
    for slide, (n, sec, label) in zip(prs.slides, PLAN):
        elapsed += sec
        stamp = (f"[{sec}s · {label} · cumulative {elapsed//60}:{elapsed%60:02d} of 20:00]\n\n")
        tf = slide.notes_slide.notes_text_frame
        body = tf.text
        if body.startswith("["):                      # already stamped, replace it
            body = body.split("]\n\n", 1)[-1]
        tf.text = stamp + body

    prs.save(DECK)
    print(f"stamped {len(PLAN)} slides · {total}s total · 20:00")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
